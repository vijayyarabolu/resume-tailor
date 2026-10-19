"""Builds the demo profile's base resume (resume_tailor/examples/alex-rivera/resumes/Alex_Rivera_ML_Engineer.docx).

Alex Rivera and every company here are made up. Needs python-docx (a dev-only dependency):
    pip install python-docx && python tools/make_example_resume.py
The layout is the one docs/BASE_RESUME.md describes: name, title and contact lines, then Summary, Skills,
Experience and Education, with each role's dates on a right tab stop.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "resume_tailor" / "examples" / "alex-rivera" / "resumes" / \
    "Alex_Rivera_ML_Engineer.docx"
FONT = "Arial"
RIGHT = Inches(7.3)

SUMMARY = ("Machine Learning Engineer with 6 years building forecasting, risk and retrieval systems that run in "
           "production. I work across the whole path from data pipelines to model serving and monitoring, mostly "
           "in Python on AWS. Recent work includes LLM retrieval tools for operations teams and models that cut "
           "costs in freight and healthcare.")
SKILLS = [
    ("Languages & Data", "Python, SQL, Scala, Bash, PostgreSQL"),
    ("Machine Learning", "PyTorch, scikit-learn, XGBoost, pandas, Feature Engineering, A/B Testing, Time Series"),
    ("MLOps & Cloud", "AWS, SageMaker, Docker, Kubernetes, MLflow, Airflow, CI/CD, Terraform"),
    ("GenAI", "LLMs, RAG, LangChain, Embeddings, pgvector, LLM Evaluation"),
]
JOBS = [
    ("Lumen Freight, USA", "Mar 2023 – Present", "Senior Machine Learning Engineer", [
        "Built a demand forecasting service in PyTorch and Airflow on AWS that predicts load volume for 1,200 "
        "routes, cutting empty truck miles by 14% across the regional network.",
        "Shipped a retrieval tool with LangChain and pgvector that answers dispatcher questions from 40,000 "
        "shipping documents, with a 600 question evaluation set run on every release.",
        "Moved model training to SageMaker with MLflow tracking and Docker images, taking retraining from two "
        "days to five hours so the team could release models every week.",
        "Set up drift monitoring and alerts for 9 production models on Kubernetes, catching two broken data "
        "pipelines before they reached the pricing team or customer quotes.",
    ]),
    ("Brightline Health, USA", "Jun 2020 – Feb 2023", "Machine Learning Engineer", [
        "Trained XGBoost models that flag patients at risk of missing follow-up visits, raising scheduled "
        "follow-ups by 11% across 23 clinics during a six month pilot.",
        "Rebuilt the feature pipeline in Python and SQL with tested, versioned transformations, cutting the "
        "nightly run from 3 hours to 35 minutes and ending most manual reruns.",
        "Ran A/B tests on appointment reminder timing with the product team and wrote the analysis templates "
        "analysts still use to size and read their experiments.",
        "Worked with clinicians to define alerts they would trust, and cut false alarms by 30% by adding visit "
        "history features and a calibrated decision threshold.",
    ]),
    ("Northwind Analytics, USA", "Jul 2018 – May 2020", "Data Scientist", [
        "Built churn models in scikit-learn for retail clients and explained each model's drivers in plain "
        "language so account managers could act on them in quarterly reviews.",
        "Wrote time series forecasts in pandas for store staffing at 150 locations, reducing overtime hours by "
        "8% in the first year after the forecasts went live.",
        "Automated weekly client reports with Python and SQL, saving the analytics team about 10 hours a week "
        "and removing the copy and paste errors clients used to flag.",
        "Joined point of sale data from 12 retail clients into one warehouse schema in SQL, which became the base "
        "for every model the team built after that.",
    ]),
]
EDUCATION = ("M.S. Computer Science, Cascade State University", "Sep 2016 – Jun 2018")


def run(p, text, bold=False, size=10):
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = FONT
    return r


def spacing(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0


def heading(doc, text):
    p = doc.add_paragraph()
    spacing(p, 10, 4)
    run(p, text, bold=True, size=11)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "555555")):
        bottom.set(qn(k), v)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def dated(doc, left, dates, bold=True):
    p = doc.add_paragraph()
    spacing(p, 6, 0)
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT, WD_TAB_ALIGNMENT.RIGHT)
    run(p, left, bold=bold)
    run(p, "\t" + dates)
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.6)
    sec.top_margin = sec.bottom_margin = Inches(0.55)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    p = doc.add_paragraph(); spacing(p); run(p, "Alex Rivera", bold=True, size=18)
    p = doc.add_paragraph(); spacing(p, 2); run(p, "Machine Learning Engineer", size=12)
    p = doc.add_paragraph(); spacing(p, 2)
    run(p, "Portland, OR"); run(p, " | alex@example.com | linkedin.com/in/example | github.com/example")
    heading(doc, "Summary")
    p = doc.add_paragraph(); spacing(p); run(p, SUMMARY)
    heading(doc, "Skills")
    for category, items in SKILLS:
        p = doc.add_paragraph(); spacing(p, 0, 1)
        run(p, f"{category}: ", bold=True); run(p, items)
    heading(doc, "Experience")
    for company, dates, title, bullets in JOBS:
        dated(doc, company, dates)
        p = doc.add_paragraph(); spacing(p, 0, 1); run(p, title).italic = True
        for b in bullets:
            p = doc.add_paragraph(style="List Bullet"); spacing(p, 0, 2); run(p, b)
    heading(doc, "Education")
    dated(doc, *EDUCATION, bold=False)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.author = "Alex Rivera"
    doc.core_properties.last_modified_by = "Alex Rivera"
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
